from pathlib import Path
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BASE_DIR = Path(__file__).parent
OUTPUT = BASE_DIR / "Multi_Agent_Research_Presentation.pptx"
RUNS_FILE = BASE_DIR / "runs.json"

# Theme
NAVY = RGBColor(8, 17, 34)
BLUE = RGBColor(29, 78, 216)
PURPLE = RGBColor(124, 58, 237)
TEAL = RGBColor(15, 118, 110)
SKY = RGBColor(219, 234, 254)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LIGHT_BG = RGBColor(245, 248, 255)
CARD = RGBColor(255, 255, 255)
SOFT_BLUE = RGBColor(232, 240, 255)
SOFT_GREEN = RGBColor(236, 253, 243)
SOFT_PURPLE = RGBColor(243, 232, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def load_metrics():
    if not RUNS_FILE.exists():
        return {
            "total_runs": 0,
            "completed_runs": 0,
            "failed_runs": 0,
            "task_success_rate": "0.0%",
            "reasoning_consistency": "0.0%",
            "error_recovery": "0.0%",
            "efficiency": "No completed runs yet",
        }
    try:
        runs = json.loads(RUNS_FILE.read_text(encoding="utf-8"))
    except Exception:
        runs = []

    total_runs = len(runs)
    completed_runs = [r for r in runs if r.get("status") == "completed"]
    failed_runs = [r for r in runs if r.get("status") == "failed"]
    success_rate = (len(completed_runs) / total_runs * 100) if total_runs else 0
    consistency_rate = (
        len([r for r in completed_runs if (r.get("result_preview") or "").strip()]) / len(completed_runs) * 100
        if completed_runs else 0
    )
    durations = [int(r.get("duration_ms", 0)) for r in completed_runs if str(r.get("duration_ms", "")).isdigit()]
    avg_runtime = f"{round(sum(durations) / len(durations), 2)} ms average runtime" if durations else "No completed runs yet"
    return {
        "total_runs": total_runs,
        "completed_runs": len(completed_runs),
        "failed_runs": len(failed_runs),
        "task_success_rate": f"{success_rate:.1f}%",
        "reasoning_consistency": f"{consistency_rate:.1f}%",
        "error_recovery": "100.0%" if total_runs and not failed_runs else (f"{(max(total_runs - len(failed_runs), 0) / total_runs * 100):.1f}%" if total_runs else "0.0%"),
        "efficiency": avg_runtime,
    }


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None, dark=False):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12.0), Inches(1.1))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = DARK if dark else WHITE

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.28), Inches(11.7), Inches(0.6))
        p2 = sub_box.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(11)
        r2.font.color.rgb = MUTED if dark else SKY


def add_footer(slide, text, dark=False):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(7.0), Inches(12.0), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED if dark else SKY


def add_bullet_list(slide, items, left, top, width, height, color=DARK, font_size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(10)


def add_card(slide, left, top, width, height, title, body, fill_color=CARD, title_color=DARK, body_color=MUTED):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = SOFT_BLUE

    title_box = slide.shapes.add_textbox(left + Inches(0.16), top + Inches(0.12), width - Inches(0.3), Inches(0.35))
    p1 = title_box.text_frame.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = title_color

    body_box = slide.shapes.add_textbox(left + Inches(0.16), top + Inches(0.46), width - Inches(0.3), height - Inches(0.55))
    tf = body_box.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(body if isinstance(body, list) else [body]):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.color.rgb = body_color
        p.space_after = Pt(6)


metrics = load_metrics()

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.28))
accent.fill.solid(); accent.fill.fore_color.rgb = BLUE; accent.line.fill.background()
add_title(
    slide,
    "Design and Evaluation of a Tool-Using Autonomous Multi-Agent AI System",
    "FastAPI + CrewAI + Firebase | Goal-Directed Problem Solving Project",
)
add_bullet_list(
    slide,
    [
        "Research project focused on reliable AI reasoning with tools",
        "Supports planning, calculations, Python analysis, budgeting, and evaluation",
        "Interactive web dashboard with sign-in, metrics, and runtime history",
    ],
    Inches(0.8), Inches(2.0), Inches(7.2), Inches(2.2), color=WHITE, font_size=18,
)
add_card(slide, Inches(8.55), Inches(1.95), Inches(3.8), Inches(2.55), "Prepared by", ["Abdulrazig Mohammed", "Member 1", "Member 2"], fill_color=RGBColor(18, 33, 63), title_color=WHITE, body_color=SKY)
add_card(slide, Inches(8.55), Inches(4.8), Inches(3.8), Inches(1.2), "Core Idea", "Plan → Act → Evaluate → Reflect", fill_color=RGBColor(15, 50, 96), title_color=WHITE, body_color=WHITE)
add_footer(slide, "Multi-Agent Research System", dark=False)

# Slide 2: Problem + motivation
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, LIGHT_BG)
add_title(slide, "Problem Statement & Motivation", "Why use multiple agents instead of a single agent?", dark=True)
add_card(slide, Inches(0.75), Inches(1.8), Inches(4.0), Inches(3.9), "Research Question", [
    "Can structured multi-agent reasoning with tool use improve reliability over a standard single-agent approach?",
    "This project measures success rate, consistency, recovery, and efficiency.",
], fill_color=CARD)
add_card(slide, Inches(4.95), Inches(1.8), Inches(3.7), Inches(3.9), "Target Tasks", [
    "• Geometry and arithmetic",
    "• Budget planning",
    "• Python data analysis",
    "• Multi-step decision support",
], fill_color=SOFT_BLUE)
add_card(slide, Inches(8.9), Inches(1.8), Inches(3.7), Inches(3.9), "Why It Matters", [
    "• Better task decomposition",
    "• Clearer validation before answers",
    "• Improved recovery from weak reasoning",
    "• Useful for demos and research evaluation",
], fill_color=SOFT_PURPLE)
add_footer(slide, "Slide 2", dark=True)

# Slide 3: System architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, LIGHT_BG)
add_title(slide, "System Architecture", "How the project is organized end to end", dark=True)
add_card(slide, Inches(0.7), Inches(1.9), Inches(2.2), Inches(1.25), "Frontend", ["Jinja2 UI", "Protected dashboard", "Auth pages"], fill_color=SOFT_BLUE)
add_card(slide, Inches(3.1), Inches(1.9), Inches(2.2), Inches(1.25), "Backend", ["FastAPI routes", "/solve, /runs", "/metrics, /health"], fill_color=SOFT_GREEN)
add_card(slide, Inches(5.5), Inches(1.9), Inches(2.2), Inches(1.25), "Agent Layer", ["Planner", "Tool Agent", "Evaluator", "Reflection"], fill_color=SOFT_PURPLE)
add_card(slide, Inches(7.9), Inches(1.9), Inches(2.2), Inches(1.25), "Tools", ["Calculator", "Python executor", "Data analysis"], fill_color=SOFT_BLUE)
add_card(slide, Inches(10.3), Inches(1.9), Inches(2.2), Inches(1.25), "Storage", ["runs.json", "Metrics", "History"], fill_color=SOFT_GREEN)

flow = slide.shapes.add_textbox(Inches(0.85), Inches(3.65), Inches(11.9), Inches(2.2))
flow_tf = flow.text_frame
for idx, line in enumerate([
    "User signs in and submits a problem from the dashboard.",
    "Fast-path logic solves simple tasks instantly when possible.",
    "CrewAI agents collaborate for more complex tasks using a structured workflow.",
    "Results, timings, and evaluation metrics are displayed in the app.",
]):
    p = flow_tf.paragraphs[0] if idx == 0 else flow_tf.add_paragraph()
    p.text = f"• {line}"
    p.font.size = Pt(17)
    p.font.color.rgb = DARK
    p.space_after = Pt(10)
add_footer(slide, "Slide 3", dark=True)

# Slide 4: Workflow
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, "Agent Workflow", "Plan → Act → Evaluate → Reflect", dark=False)
steps = [
    ("1. Planner", "Breaks the task into a clear execution strategy."),
    ("2. Tool Agent", "Uses calculations, Python, or structured reasoning to solve it."),
    ("3. Evaluator", "Checks correctness, consistency, and completeness."),
    ("4. Reflection", "Suggests improvements for future reliability and recovery."),
]
for i, (title, body) in enumerate(steps):
    add_card(slide, Inches(0.8 + (i % 2) * 6.1), Inches(1.9 + (i // 2) * 2.05), Inches(5.6), Inches(1.6), title, body, fill_color=RGBColor(18, 33, 63), title_color=WHITE, body_color=SKY)
add_footer(slide, "Slide 4", dark=False)

# Slide 5: Features and UI
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, LIGHT_BG)
add_title(slide, "User Experience & Key Features", "The app combines research functionality with a polished interface", dark=True)
add_card(slide, Inches(0.7), Inches(1.8), Inches(3.9), Inches(1.65), "Authentication", ["Firebase sign-up / sign-in", "Protected `/app` dashboard"], fill_color=SOFT_BLUE)
add_card(slide, Inches(4.75), Inches(1.8), Inches(3.9), Inches(1.65), "Interactive Dashboard", ["Prompt input", "AI response panel", "Recent activity and runtime chart"], fill_color=SOFT_GREEN)
add_card(slide, Inches(8.8), Inches(1.8), Inches(3.9), Inches(1.65), "Fast Problem Solving", ["Geometry", "Budgets", "Planning", "Data analysis"], fill_color=SOFT_PURPLE)
add_bullet_list(
    slide,
    [
        "Modern visual design with chart-driven evaluation overview",
        "One-screen layout suitable for demos and presentations",
        "Run history and metrics update as tasks are completed",
        "Browser UI plus API access for testing and deployment",
    ],
    Inches(0.9), Inches(4.0), Inches(11.6), Inches(2.0), color=DARK, font_size=17,
)
add_footer(slide, "Slide 5", dark=True)

# Slide 6: Metrics
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, LIGHT_BG)
add_title(slide, "Evaluation Metrics", "Live metrics generated from recorded runs", dark=True)
metric_items = [
    ("Task Success Rate", metrics["task_success_rate"], SOFT_BLUE),
    ("Reasoning Consistency", metrics["reasoning_consistency"], SOFT_GREEN),
    ("Error Recovery", metrics["error_recovery"], SOFT_PURPLE),
    ("Efficiency", metrics["efficiency"], SOFT_BLUE),
]
positions = [(0.8, 1.9), (6.7, 1.9), (0.8, 3.8), (6.7, 3.8)]
for (label, value, color), (x, y) in zip(metric_items, positions):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.0), Inches(1.45))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.color.rgb = SOFT_BLUE
    tb = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.14), Inches(4.6), Inches(1.0))
    tf = tb.text_frame
    p1 = tf.paragraphs[0]
    r1 = p1.add_run(); r1.text = label; r1.font.size = Pt(14); r1.font.bold = True; r1.font.color.rgb = DARK
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = value; r2.font.size = Pt(18); r2.font.bold = True; r2.font.color.rgb = BLUE

add_card(slide, Inches(0.8), Inches(5.7), Inches(10.9), Inches(0.8), "Current dataset snapshot", f"Total runs: {metrics['total_runs']} | Completed runs: {metrics['completed_runs']} | Failed runs: {metrics['failed_runs']}", fill_color=CARD)
add_footer(slide, "Slide 6", dark=True)

# Slide 7: Tech stack and deployment
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, LIGHT_BG)
add_title(slide, "Technology Stack & Deployment", "Built for local testing and cloud hosting", dark=True)
add_card(slide, Inches(0.75), Inches(1.8), Inches(4.0), Inches(3.6), "Backend Stack", [
    "• FastAPI + Uvicorn",
    "• CrewAI orchestration",
    "• Pydantic validation",
    "• python-dotenv configuration",
], fill_color=SOFT_BLUE)
add_card(slide, Inches(4.95), Inches(1.8), Inches(4.0), Inches(3.6), "Frontend + Auth", [
    "• Jinja2 templates",
    "• Styled dashboard UI",
    "• Firebase Email/Password auth",
    "• Live metrics and run history",
], fill_color=SOFT_GREEN)
add_card(slide, Inches(9.15), Inches(1.8), Inches(3.45), Inches(3.6), "Deployment", [
    "• Dockerfile included",
    "• Render deployment ready",
    "• /health for monitoring",
    "• API docs at /docs",
], fill_color=SOFT_PURPLE)
add_footer(slide, "Slide 7", dark=True)

# Slide 8: Closing
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, "Conclusion & Future Improvements", "What this project demonstrates", dark=False)
add_bullet_list(
    slide,
    [
        "A structured multi-agent workflow can make problem solving more explainable and evaluation-friendly.",
        "Tool use plus validation helps support accurate, goal-directed answers for common tasks.",
        "The project now includes a modern dashboard, authentication, metrics, and deployment-ready APIs.",
        "Future work: stronger server-side auth, deeper benchmarking, and larger evaluation datasets.",
    ],
    Inches(0.8), Inches(1.9), Inches(11.7), Inches(3.2), color=WHITE, font_size=18,
)
add_card(slide, Inches(0.9), Inches(5.4), Inches(11.5), Inches(0.95), "Thank you", "Questions, demo, and walkthrough", fill_color=RGBColor(18, 33, 63), title_color=WHITE, body_color=WHITE)
add_footer(slide, "Multi-Agent Research System", dark=False)

prs.save(OUTPUT)
print(f"PRESENTATION_CREATED: {OUTPUT}")
